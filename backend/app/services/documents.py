"""The document centre: every document an admin can download (Excel, Word, PDF) or get in the bot.

``TEMPLATES`` is the list the admin panel shows («📄 Hujjatlar»); each template says which
permission it needs and which filters it takes. Builders read the database with a few grouped
queries (no query per member), so they stay fast with thousands of members.
"""

from __future__ import annotations

import logging
import uuid
from collections import defaultdict
from dataclasses import dataclass
from datetime import UTC, date, datetime, time, timedelta
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import BadRequestError, NotFoundError
from app.models import (
    AttendanceRecord,
    Book,
    BookRead,
    Certificate,
    Event,
    Faculty,
    ScoreTransaction,
    TestAttempt,
    User,
    UserAchievement,
)
from app.models.enums import AttemptStatus, AttendanceStatus, ReadingStatus, ScoreCategory, UserStatus
from app.security.permissions import Perm
from app.services.document_render import Document, QrEntry, Table
from app.services.settings import get_setting
from app.utils.timezone import format_date, format_datetime, local_tz

logger = logging.getLogger(__name__)

FINISHED = (AttemptStatus.SUBMITTED, AttemptStatus.EXPIRED)
MAX_APPENDIX = 400  # QR cards in one appendix (a few dozen pages)


@dataclass(frozen=True)
class Template:
    code: str
    title: str
    description: str
    formats: tuple[str, ...]
    permission: Perm
    params: tuple[str, ...]  # period, faculty, course, group, event, test, report


TEMPLATES: tuple[Template, ...] = (
    Template(
        "students",
        "Talabalar faoliyati",
        "Har bir talaba: tadbirlar, testlar, sovrinli o'rinlar, o'qilgan kitoblar, sertifikatlar, yutuqlar va "
        "ballar — fakultet, kurs, guruh bo'yicha. Institut reyting tizimi uchun.",
        ("xlsx", "pdf", "docx"),
        Perm.ANALYTICS_VIEW,
        ("period", "faculty", "course", "group"),
    ),
    Template(
        "statistics",
        "Statistik hisobot + QR kodli ilova",
        "Fakultet, kurs va guruhlar kesimida a'zolar va faollik soni; ilovada faol a'zolar QR kodi bilan "
        "(skanerlansa — a'zoning klubdagi faoliyati ochiladi).",
        ("pdf", "docx", "xlsx"),
        Perm.ANALYTICS_VIEW,
        ("period", "faculty", "course", "group"),
    ),
    Template(
        "attendance",
        "Tadbir davomati ro'yxati",
        "Tadbirda qatnashganlar: fakultet, kurs, guruh, smena, qayd usuli va vaqti.",
        ("xlsx", "pdf", "docx"),
        Perm.ATTENDANCE_VIEW,
        ("event",),
    ),
    Template(
        "test_results",
        "Test natijalari",
        "O'rinlar, ballar va vaqt; tadbir testida davomat qilmaganlar «davomatsiz» belgisi bilan.",
        ("xlsx", "pdf", "docx"),
        Perm.TEST_VIEW,
        ("test",),
    ),
    Template(
        "members",
        "A'zolar ro'yxati",
        "Barcha a'zolar: fakultet, kurs, guruh, a'zo raqami, holati va qo'shilgan sanasi.",
        ("xlsx", "pdf", "docx"),
        Perm.USER_VIEW,
        ("faculty", "course", "group"),
    ),
    Template(
        "certificates",
        "Sertifikatlar reyestri",
        "Davr ichida berilgan sertifikatlar va ularning tekshirish kodlari.",
        ("xlsx", "pdf", "docx"),
        Perm.ANALYTICS_VIEW,
        ("period",),
    ),
    Template(
        "books",
        "Kitoblar va o'quvchilar",
        "Har bir kitob: o'qiyotganlar, davr ichida o'qib tugatganlar va tasdiqlanganlar soni.",
        ("xlsx", "pdf", "docx"),
        Perm.BOOK_VIEW,
        ("period",),
    ),
    Template(
        "event_report",
        "Tadbir hisoboti (rasmiy)",
        "«Hisobotlar» bo'limida tayyorlangan hisobot: tasdiqlangan rasmiy PDF yoki tahrirlash uchun Word.",
        ("pdf", "docx"),
        Perm.REPORT_VIEW,
        ("report",),
    ),
)
BY_CODE = {template.code: template for template in TEMPLATES}


@dataclass
class Params:
    date_from: date | None = None
    date_to: date | None = None
    faculty_id: uuid.UUID | None = None
    course: int | None = None
    group_name: str | None = None
    event_id: uuid.UUID | None = None
    test_id: uuid.UUID | None = None
    report_id: uuid.UUID | None = None


def academic_year(today: date | None = None) -> tuple[date, date]:
    """1 September – 31 August of the current study year."""
    today = today or datetime.now(local_tz()).date()
    start_year = today.year if today.month >= 9 else today.year - 1
    return date(start_year, 9, 1), date(start_year + 1, 8, 31)


def _period(params: Params) -> tuple[date, date, datetime, datetime]:
    default_start, default_end = academic_year()
    start = params.date_from or default_start
    end = params.date_to or default_end
    if end < start:
        raise BadRequestError("PERIOD_INVALID", "Davr oxiri boshidan oldin bo'lmasin.")
    tz = local_tz()
    return start, end, datetime.combine(start, time.min, tz), datetime.combine(end + timedelta(days=1), time.min, tz)


def _day(value: date) -> str:
    return value.strftime("%d.%m.%Y")


async def _header(session: AsyncSession) -> list[str]:
    organization = str(await get_setting(session, "org.organization_name"))
    club = str(await get_setting(session, "org.club_name"))
    return [line for line in (organization, club) if line]


async def _signature(session: AsyncSession) -> str:
    title = str(await get_setting(session, "certificates.signer_title") or "Klub rahbari")
    name = str(await get_setting(session, "certificates.signer_name") or "")
    return f"{title}: ____________________ {name}".rstrip()


def _member_filter(params: Params) -> list[Any]:
    from app.services.members import normalize_group

    conditions: list[Any] = [User.deleted_at.is_(None), User.status != UserStatus.PENDING]
    if params.faculty_id is not None:
        conditions.append(User.faculty_id == params.faculty_id)
    if params.course is not None:
        conditions.append(User.course == params.course)
    if params.group_name:
        conditions.append(User.group_name == normalize_group(params.group_name))
    return conditions


async def _filter_lines(session: AsyncSession, params: Params) -> list[str]:
    parts = []
    if params.faculty_id is not None:
        faculty = await session.get(Faculty, params.faculty_id)
        parts.append(f"Fakultet: {faculty.name if faculty else '—'}")
    if params.course is not None:
        parts.append(f"Kurs: {params.course}")
    if params.group_name:
        parts.append(f"Guruh: {params.group_name}")
    return ["; ".join(parts)] if parts else []


# ---------------------------------------------------------------------------- activity of members


async def _activity(session: AsyncSession, params: Params) -> tuple[list[tuple[User, str | None]], dict[str, Any]]:
    """Members (with the faculty name) and their activity in the period, counted in a few grouped queries."""
    _, _, since, until = _period(params)
    members = [
        (user, faculty)
        for user, faculty in (
            await session.execute(
                select(User, Faculty.name)
                .outerjoin(Faculty, Faculty.id == User.faculty_id)
                .where(*_member_filter(params))
                .order_by(Faculty.sort_order.nulls_last(), Faculty.name, User.course, User.group_name, User.last_name)
            )
        ).all()
    ]
    ids = [user.id for user, _ in members]

    async def counts(query: Any) -> dict[uuid.UUID, int]:
        return {user_id: int(value or 0) for user_id, value in (await session.execute(query)).all()}

    events = await counts(
        select(AttendanceRecord.user_id, func.count())
        .join(Event, Event.id == AttendanceRecord.event_id)
        .where(
            AttendanceRecord.user_id.in_(ids),
            AttendanceRecord.status.in_((AttendanceStatus.VERIFIED, AttendanceStatus.SUSPICIOUS)),
            Event.starts_at >= since,
            Event.starts_at < until,
        )
        .group_by(AttendanceRecord.user_id)
    )
    best = (
        select(TestAttempt.user_id, TestAttempt.test_id, func.max(TestAttempt.score).label("score"))
        .where(
            TestAttempt.user_id.in_(ids),
            TestAttempt.status.in_(FINISHED),
            TestAttempt.submitted_at >= since,
            TestAttempt.submitted_at < until,
        )
        .group_by(TestAttempt.user_id, TestAttempt.test_id)
        .subquery()
    )
    tests: dict[uuid.UUID, int] = {}
    test_points: dict[uuid.UUID, int] = {}
    for user_id, number, total in (
        await session.execute(
            select(best.c.user_id, func.count(), func.coalesce(func.sum(best.c.score), 0)).group_by(best.c.user_id)
        )
    ).all():
        tests[user_id], test_points[user_id] = int(number), int(total)
    places = await counts(
        select(ScoreTransaction.user_id, func.count())
        .where(
            ScoreTransaction.user_id.in_(ids),
            ScoreTransaction.category == ScoreCategory.TEST_RANK_BONUS,
            ScoreTransaction.amount > 0,
            ScoreTransaction.created_at >= since,
            ScoreTransaction.created_at < until,
        )
        .group_by(ScoreTransaction.user_id)
    )
    books = await counts(
        select(BookRead.user_id, func.count())
        .where(
            BookRead.user_id.in_(ids),
            BookRead.status == ReadingStatus.COMPLETED,
            BookRead.completed_at >= since,
            BookRead.completed_at < until,
        )
        .group_by(BookRead.user_id)
    )
    certificates = await counts(
        select(Certificate.user_id, func.count())
        .where(
            Certificate.user_id.in_(ids),
            Certificate.revoked_at.is_(None),
            Certificate.issued_at >= since,
            Certificate.issued_at < until,
        )
        .group_by(Certificate.user_id)
    )
    achievements = await counts(
        select(UserAchievement.user_id, func.count())
        .where(
            UserAchievement.user_id.in_(ids),
            UserAchievement.unlocked_at >= since,
            UserAchievement.unlocked_at < until,
        )
        .group_by(UserAchievement.user_id)
    )
    points = await counts(
        select(ScoreTransaction.user_id, func.coalesce(func.sum(ScoreTransaction.amount), 0))
        .where(
            ScoreTransaction.user_id.in_(ids),
            ScoreTransaction.created_at >= since,
            ScoreTransaction.created_at < until,
        )
        .group_by(ScoreTransaction.user_id)
    )
    data = {
        "events": events,
        "tests": tests,
        "test_points": test_points,
        "places": places,
        "books": books,
        "certificates": certificates,
        "achievements": achievements,
        "points": points,
    }
    return members, data


def _active(user_id: uuid.UUID, data: dict[str, Any]) -> bool:
    return any(data[key].get(user_id) for key in ("events", "tests", "books", "certificates"))


def _member_url(user: User) -> str:
    from app.services import passport

    return passport.verify_url(passport.verify_code(user)) or f"Kitobxon Hub · {passport.member_number(user)}"


async def students(session: AsyncSession, params: Params) -> Document:
    from app.services.passport import member_number

    start, end, _, _ = _period(params)
    members, data = await _activity(session, params)
    rows = []
    totals = [0] * 8
    keys = ("events", "tests", "test_points", "places", "books", "certificates", "achievements", "points")
    for user, faculty in members:
        if user.status != UserStatus.ACTIVE and not _active(user.id, data):
            continue
        values = [data[key].get(user.id, 0) for key in keys]
        totals = [a + b for a, b in zip(totals, values, strict=True)]
        rows.append(
            [
                len(rows) + 1,
                f"{user.last_name} {user.first_name}",
                faculty or "—",
                user.course or "—",
                user.group_name or "—",
                member_number(user),
                *values,
            ]
        )
    return Document(
        title="Talabalarning kitobxonlik faoliyati",
        filename=f"talabalar-faoliyati-{start:%Y%m%d}-{end:%Y%m%d}",
        lines=[*await _header(session), f"Davr: {_day(start)} — {_day(end)}", *await _filter_lines(session, params)],
        tables=[
            Table(
                None,
                [
                    "№",
                    "F.I.Sh.",
                    "Fakultet",
                    "Kurs",
                    "Guruh",
                    "A'zo №",
                    "Tadbirlar",
                    "Testlar",
                    "Test ballari",
                    "Sovrinli o'rinlar",
                    "O'qilgan kitoblar",
                    "Sertifikatlar",
                    "Yutuqlar",
                    "Jami ball",
                ],
                rows,
                totals=["", f"Jami: {len(rows)} nafar", "", "", "", "", *totals],
            )
        ],
        landscape=True,
        signature=await _signature(session),
    )


async def statistics(session: AsyncSession, params: Params) -> Document:
    start, end, _, _ = _period(params)
    members, data = await _activity(session, params)

    def summarize(key_of: Any) -> list[list[Any]]:
        groups: dict[Any, list[int]] = defaultdict(lambda: [0] * 6)
        for user, faculty in members:
            key = key_of(user, faculty)
            row = groups[key]
            row[0] += 1
            row[1] += 1 if _active(user.id, data) else 0
            row[2] += data["events"].get(user.id, 0)
            row[3] += data["tests"].get(user.id, 0)
            row[4] += data["books"].get(user.id, 0)
            row[5] += data["points"].get(user.id, 0)
        return [[key, *values] for key, values in groups.items()]

    columns = ["A'zolar", "Faol", "Davomatlar", "Testlar", "O'qilgan kitoblar", "Ballar"]
    by_faculty = summarize(lambda user, faculty: faculty or "Fakultetsiz")
    by_course = sorted(summarize(lambda user, faculty: f"{user.course}-kurs" if user.course else "—"), key=str)
    by_group = sorted(
        summarize(lambda user, faculty: f"{user.group_name or '—'} ({faculty or '—'})"),
        key=lambda row: -row[2],
    )
    totals = ["Jami", *[sum(row[index] for row in by_faculty) for index in range(1, 7)]]
    active = [(user, faculty) for user, faculty in members if _active(user.id, data)]
    appendix = [
        QrEntry(
            title=f"{user.last_name} {user.first_name}",
            details=" · ".join(
                part for part in (faculty, f"{user.course}-kurs" if user.course else None, user.group_name) if part
            ),
            url=_member_url(user),
        )
        for user, faculty in active[:MAX_APPENDIX]
    ]
    lines = [*await _header(session), f"Davr: {_day(start)} — {_day(end)}", *await _filter_lines(session, params)]
    lines.append(f"A'zolar: {len(members)} nafar, shundan faol: {len(active)} nafar")
    if len(active) > MAX_APPENDIX:
        lines.append(f"Ilovada birinchi {MAX_APPENDIX} nafar faol a'zo (to'liq ro'yxat — «Talabalar faoliyati»)")
    return Document(
        title="Kitobxonlar klubi faoliyati bo'yicha statistik hisobot",
        filename=f"statistika-{start:%Y%m%d}-{end:%Y%m%d}",
        lines=lines,
        tables=[
            Table("Fakultetlar kesimida", ["Fakultet", *columns], by_faculty, totals=totals),
            Table("Kurslar kesimida", ["Kurs", *columns], by_course),
            Table("Guruhlar kesimida (eng faollari yuqorida)", ["Guruh", *columns], by_group[:150]),
        ],
        appendix_title="Ilova: faol a'zolar ro'yxati (QR kod bilan)",
        appendix=appendix,
        signature=await _signature(session),
    )


# ---------------------------------------------------------------------------- one event, one test


async def attendance(session: AsyncSession, params: Params) -> Document:
    from app.models import EventShift
    from app.services.event_times import shift_text

    if params.event_id is None:
        raise BadRequestError("EVENT_REQUIRED", "Tadbirni tanlang.")
    event = await session.get(Event, params.event_id)
    if event is None or event.deleted_at is not None:
        raise NotFoundError("EVENT_NOT_FOUND", "Tadbir topilmadi.")
    shifts = {
        shift.id: shift for shift in await session.scalars(select(EventShift).where(EventShift.event_id == event.id))
    }
    from app.models import EventRegistration

    registered_shift = dict(
        (
            await session.execute(
                select(EventRegistration.user_id, EventRegistration.shift_id).where(
                    EventRegistration.event_id == event.id
                )
            )
        ).all()
    )
    methods = {"QR": "QR kod", "CODE": "Kod", "MANUAL": "Admin"}
    statuses = {
        "VERIFIED": "Tasdiqlangan",
        "SUSPICIOUS": "Shubhali",
        "REVOKED": "Bekor qilingan",
        "PENDING": "Kutilmoqda",
    }
    rows = []
    for record, user, faculty in (
        await session.execute(
            select(AttendanceRecord, User, Faculty.name)
            .join(User, User.id == AttendanceRecord.user_id)
            .outerjoin(Faculty, Faculty.id == User.faculty_id)
            .where(AttendanceRecord.event_id == event.id)
            .order_by(Faculty.name, User.course, User.group_name, User.last_name)
        )
    ).all():
        shift = shifts.get(registered_shift.get(user.id)) if registered_shift.get(user.id) else None
        rows.append(
            [
                len(rows) + 1,
                f"{user.last_name} {user.first_name}",
                faculty or "—",
                user.course or "—",
                user.group_name or "—",
                shift_text(shift) if shift else "—",
                methods.get(str(record.method), str(record.method)),
                statuses.get(str(record.status), str(record.status)),
                format_datetime(record.checked_in_at),
            ]
        )
    counted = sum(1 for row in rows if row[7] in ("Tasdiqlangan", "Shubhali"))
    return Document(
        title=f"«{event.title}» — davomat ro'yxati",
        filename=f"davomat-{format_date(event.starts_at).replace('.', '-')}",
        lines=[
            *await _header(session),
            f"Sana: {format_datetime(event.starts_at)} · {event.location or ''}".rstrip(" ·"),
            f"Qatnashganlar: {counted} nafar",
        ],
        tables=[
            Table(
                None,
                ["№", "F.I.Sh.", "Fakultet", "Kurs", "Guruh", "Smena", "Usul", "Holat", "Vaqt"],
                rows,
            )
        ],
        landscape=True,
        signature=await _signature(session),
    )


async def test_results(session: AsyncSession, params: Params) -> Document:
    from app.services import quizzes

    if params.test_id is None:
        raise BadRequestError("TEST_REQUIRED", "Testni tanlang.")
    test, items = await quizzes.results(session, params.test_id)
    best: dict[uuid.UUID, dict[str, Any]] = {}
    for item in items:
        if item["score"] is None:
            continue
        current = best.get(item["user_id"])
        if current is None or item["score"] > current["score"]:
            best[item["user_id"]] = item
    ordered = sorted(best.values(), key=lambda item: (item["place"] or 10**6, -(item["score"] or 0)))
    rows = []
    for item in ordered:
        minutes = (item["duration_seconds"] or 0) // 60
        seconds = (item["duration_seconds"] or 0) % 60
        note = "davomatsiz" if item.get("attended") is False else ""
        rows.append(
            [
                item["place"] or "—",
                f"{item['last_name']} {item['first_name']}",
                item["faculty"] or "—",
                item["course"] or "—",
                item["group_name"] or "—",
                f"{item['score']}/{test.max_score}",
                f"{item['correct_count'] or 0}/{item['question_count']}",
                f"{minutes}:{seconds:02d}",
                note,
            ]
        )
    return Document(
        title=f"«{test.title}» testi natijalari",
        filename=f"test-natijalari-{format_date(test.starts_at).replace('.', '-')}",
        lines=[
            *await _header(session),
            f"{format_datetime(test.starts_at)} — {format_datetime(test.ends_at)}",
            f"Qatnashchilar: {len(rows)} nafar",
        ],
        tables=[
            Table(None, ["O'rin", "F.I.Sh.", "Fakultet", "Kurs", "Guruh", "Ball", "To'g'ri", "Vaqt", "Izoh"], rows)
        ],
        landscape=True,
        signature=await _signature(session),
    )


# ---------------------------------------------------------------------------- lists


async def members_list(session: AsyncSession, params: Params) -> Document:
    from app.services.passport import member_number

    statuses = {"ACTIVE": "Faol", "INACTIVE": "Faol emas", "BLOCKED": "Bloklangan", "LEFT": "Chiqqan"}
    rows = []
    for user, faculty in (
        await session.execute(
            select(User, Faculty.name)
            .outerjoin(Faculty, Faculty.id == User.faculty_id)
            .where(*_member_filter(params))
            .order_by(Faculty.name, User.course, User.group_name, User.last_name)
        )
    ).all():
        rows.append(
            [
                len(rows) + 1,
                f"{user.last_name} {user.first_name}",
                faculty or "—",
                user.course or "—",
                user.group_name or "—",
                member_number(user),
                statuses.get(str(user.status), str(user.status)),
                format_date(user.registered_at or user.created_at),
            ]
        )
    return Document(
        title="Kitobxonlar klubi a'zolari ro'yxati",
        filename=f"azolar-{datetime.now(UTC):%Y%m%d}",
        lines=[
            *await _header(session),
            f"Sana: {format_date(datetime.now(UTC))}",
            *await _filter_lines(session, params),
        ],
        tables=[Table(None, ["№", "F.I.Sh.", "Fakultet", "Kurs", "Guruh", "A'zo №", "Holati", "Qo'shilgan"], rows)],
        landscape=True,
    )


async def certificates_list(session: AsyncSession, params: Params) -> Document:
    from app.services.certificates import display_code

    start, end, since, until = _period(params)
    kinds = {"TEST_PLACE": "Test g'olibi", "MONTHLY_TOP": "Oy yetakchisi", "MANUAL": "Qo'lda berilgan"}
    rows = []
    for certificate, user in (
        await session.execute(
            select(Certificate, User)
            .join(User, User.id == Certificate.user_id)
            .where(Certificate.issued_at >= since, Certificate.issued_at < until)
            .order_by(Certificate.issued_at)
        )
    ).all():
        rows.append(
            [
                len(rows) + 1,
                f"{user.last_name} {user.first_name}",
                user.group_name or "—",
                certificate.title,
                kinds.get(str(certificate.kind), str(certificate.kind)),
                format_date(certificate.issued_at),
                display_code(certificate.verify_code),
                "bekor qilingan" if certificate.revoked_at else "",
            ]
        )
    return Document(
        title="Sertifikatlar reyestri",
        filename=f"sertifikatlar-{start:%Y%m%d}-{end:%Y%m%d}",
        lines=[*await _header(session), f"Davr: {_day(start)} — {_day(end)}", f"Jami: {len(rows)} ta"],
        tables=[Table(None, ["№", "F.I.Sh.", "Guruh", "Sertifikat", "Turi", "Sana", "Kod", "Izoh"], rows)],
        landscape=True,
    )


async def books_list(session: AsyncSession, params: Params) -> Document:
    start, end, since, until = _period(params)
    readers = dict((await session.execute(select(BookRead.book_id, func.count()).group_by(BookRead.book_id))).all())
    finished = dict(
        (
            await session.execute(
                select(BookRead.book_id, func.count())
                .where(
                    BookRead.status == ReadingStatus.COMPLETED,
                    BookRead.completed_at >= since,
                    BookRead.completed_at < until,
                )
                .group_by(BookRead.book_id)
            )
        ).all()
    )
    verified = dict(
        (
            await session.execute(
                select(BookRead.book_id, func.count())
                .where(BookRead.verified_at >= since, BookRead.verified_at < until)
                .group_by(BookRead.book_id)
            )
        ).all()
    )
    rows = []
    for book in await session.scalars(select(Book).where(Book.deleted_at.is_(None)).order_by(Book.created_at)):
        rows.append(
            [
                len(rows) + 1,
                book.title,
                book.author,
                _day(book.deadline) if book.deadline else "—",
                int(readers.get(book.id, 0)),
                int(finished.get(book.id, 0)),
                int(verified.get(book.id, 0)),
            ]
        )
    return Document(
        title="Kitoblar va o'quvchilar",
        filename=f"kitoblar-{start:%Y%m%d}-{end:%Y%m%d}",
        lines=[*await _header(session), f"Davr: {_day(start)} — {_day(end)}"],
        tables=[
            Table(None, ["№", "Kitob", "Muallif", "Muddat", "O'quvchilar", "Tugatganlar", "Tasdiqlanganlar"], rows)
        ],
    )


# ---------------------------------------------------------------------------- the official event report in Word


async def report_docx(session: AsyncSession, report_id: uuid.UUID) -> tuple[str, bytes]:
    """The event report as an editable Word file: the same text, facts and photos as the PDF."""
    import asyncio
    import io

    from docx import Document as WordDocument
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    from app.core.errors import AppError
    from app.services import reports
    from app.services.report_pdf import body_blocks
    from app.services.storage import get_storage

    report = await reports.get_report(session, report_id)
    event = await session.get(Event, report.event_id) if report.event_id else None
    photos = []
    for photo in await reports.list_photos(session, report.id):
        try:
            photos.append((await get_storage().get(photo.object_key), photo.caption))
        except AppError:
            continue
    header = await _header(session)
    facts = report.facts_snapshot or {}

    def build() -> bytes:
        word = WordDocument()
        section = word.sections[0]
        for side in ("left_margin", "right_margin"):
            setattr(section, side, Cm(2.2))
        style = word.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        for line in header:
            paragraph = word.add_paragraph(line)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title = word.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(report.title or reports.subject(report, event.title if event else None))
        run.bold = True
        run.font.size = Pt(15)
        about = facts.get("event") or {}
        attendance = facts.get("attendance") or {}
        book = f"«{about.get('book_title')}» — {about.get('book_author')}" if about.get("book_title") else None
        summary = [
            ("Tadbir", about.get("title")),
            ("Sana va vaqt", " ".join(str(part) for part in (about.get("date"), about.get("time")) if part)),
            ("Joy", about.get("location")),
            ("Kitob", book),
            ("Mas'ul", about.get("responsible")),
            ("Ro'yxatdan o'tganlar", attendance.get("registered")),
            ("Qatnashganlar", attendance.get("attended")),
        ]
        rows = [(label, value) for label, value in summary if value not in (None, "")]
        if rows:
            grid = word.add_table(rows=0, cols=2)
            grid.style = "Table Grid"
            for label, value in rows:
                cells = grid.add_row().cells
                cells[0].text = label
                cells[1].text = str(value)
            word.add_paragraph()
        for block in body_blocks(report.body_text):
            if block["kind"] == "heading":
                heading = word.add_paragraph()
                heading_run = heading.add_run(block["text"])
                heading_run.bold = True
            else:
                paragraph = word.add_paragraph(block["text"])
                paragraph.paragraph_format.first_line_indent = Cm(1.25)
        for data, caption in photos:
            try:
                word.add_picture(io.BytesIO(data), width=Cm(15))
            except Exception:  # an unusual image format: skip it, keep the text
                logger.info("A report photo could not be put into the Word file")
                continue
            if caption:
                note = word.add_paragraph(caption)
                note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        word.add_paragraph()
        output = io.BytesIO()
        word.save(output)
        return output.getvalue()

    data = await asyncio.to_thread(build)
    return f"hisobot-{str(report.id)[:8]}", data


# ---------------------------------------------------------------------------- a member's own documents


async def portfolio_document(session: AsyncSession, user: User) -> Document:
    """The reader's portfolio as a Word-ready document (the PDF has its own design)."""
    from app.services import portfolio

    data = await portfolio.collect(session, user)
    lines = [
        data["organization"],
        data["club"],
        f"{data['name']} · {data['number']}",
        " · ".join(
            part
            for part in (data["faculty"], f"{data['course']}-kurs" if data["course"] else None, data["group"])
            if part
        ),
        f"A'zo: {data['since']} dan · Jami ball: {data['points']}"
        + (f" · Umumiy reytingda {data['rank']}-o'rin" if data["rank"] else ""),
    ]
    tables = [
        Table(
            "📚 O'qilgan kitoblar",
            ["№", "Kitob", "Muallif", "Sana", "Tasdiq"],
            [
                [index, book["title"], book["author"], book["date"], "tasdiqlangan" if book["verified"] else ""]
                for index, book in enumerate(data["books"], 1)
            ],
        ),
        Table(
            "📅 Tadbirlar",
            ["№", "Tadbir", "Sana"],
            [[index, event["title"], event["date"]] for index, event in enumerate(data["events"], 1)],
        ),
        Table(
            "📝 Testlar",
            ["№", "Test", "Sana", "Ball", "O'rin"],
            [
                [index, test["title"], test["date"], f"{test['score']}/{test['max']}", test["place"] or ""]
                for index, test in enumerate(data["tests"], 1)
            ],
        ),
        Table(
            "🎓 Sertifikatlar",
            ["№", "Sertifikat", "Sana", "Tekshirish kodi"],
            [[index, item["title"], item["date"], item["code"]] for index, item in enumerate(data["certificates"], 1)],
        ),
        Table(
            "🏅 Yutuqlar",
            ["№", "Yutuq", "Izoh", "Sana"],
            [
                [index, item["title"], item["description"], item["date"]]
                for index, item in enumerate(data["achievements"], 1)
            ],
        ),
    ]
    appendix = (
        [QrEntry(title=data["name"], details="Faoliyatni tekshirish", url=data["verify"])] if data["verify"] else []
    )
    return Document(
        title="Kitobxonlik portfoliosi",
        filename=portfolio.filename(user).removesuffix(".pdf"),
        lines=[line for line in lines if line],
        tables=[table for table in tables if table.rows],
        appendix_title="Tekshirish uchun QR kod",
        appendix=appendix,
    )


async def statement(session: AsyncSession, user: User) -> Document:
    """«Faoliyat ma'lumotnomasi»: what the member did in this study year, for the institute."""
    from app.services.passport import member_number

    start, end = academic_year()
    members, data = await _activity_for(session, user, Params(date_from=start, date_to=end))
    faculty = members[0][1] if members else None
    rows = [
        ["Tadbirlarda qatnashgan", data["events"].get(user.id, 0)],
        ["Yechgan testlari", data["tests"].get(user.id, 0)],
        ["Testlardagi ballari", data["test_points"].get(user.id, 0)],
        ["Sovrinli o'rinlar (1–3)", data["places"].get(user.id, 0)],
        ["O'qib tugatgan kitoblari", data["books"].get(user.id, 0)],
        ["Olgan sertifikatlari", data["certificates"].get(user.id, 0)],
        ["Yutuqlari", data["achievements"].get(user.id, 0)],
        ["To'plagan ballari", data["points"].get(user.id, 0)],
    ]
    url = _member_url(user)
    return Document(
        title="Ma'lumotnoma",
        filename=f"malumotnoma-{member_number(user)}",
        lines=[
            *await _header(session),
            f"{user.last_name} {user.first_name} ({member_number(user)})",
            " · ".join(
                part for part in (faculty, f"{user.course}-kurs" if user.course else None, user.group_name) if part
            ),
            f"{_day(start)} — {_day(end)} o'quv yilida kitobxonlar klubidagi faoliyati haqida",
        ],
        tables=[Table(None, ["Ko'rsatkich", "Soni"], rows)],
        appendix_title="Tekshirish uchun QR kod",
        appendix=[QrEntry(title=f"{user.last_name} {user.first_name}", details=member_number(user), url=url)],
        signature=await _signature(session),
    )


async def _activity_for(
    session: AsyncSession, user: User, params: Params
) -> tuple[list[tuple[User, str | None]], dict[str, Any]]:
    """The same counts as for the whole club, for one member."""
    faculty_params = Params(
        date_from=params.date_from,
        date_to=params.date_to,
        faculty_id=user.faculty_id,
        course=user.course,
        group_name=user.group_name,
    )
    members, data = await _activity(session, faculty_params)
    return [(member, faculty) for member, faculty in members if member.id == user.id], data


BUILDERS = {
    "students": students,
    "statistics": statistics,
    "attendance": attendance,
    "test_results": test_results,
    "members": members_list,
    "certificates": certificates_list,
    "books": books_list,
}


async def produce(session: AsyncSession, code: str, fmt: str, params: Params) -> tuple[str, bytes]:
    """(file name, file) of a template in the asked format."""
    from app.services import document_render

    template = BY_CODE.get(code)
    if template is None:
        raise NotFoundError("DOCUMENT_NOT_FOUND", "Bunday hujjat yo'q.")
    if fmt not in template.formats:
        raise BadRequestError("FORMAT_UNSUPPORTED", "Bu hujjat bunday formatda tayyorlanmaydi.")
    if code == "event_report":
        if params.report_id is None:
            raise BadRequestError("REPORT_REQUIRED", "Hisobotni tanlang.")
        if fmt == "docx":
            name, data = await report_docx(session, params.report_id)
            return f"{name}.docx", data
        from app.services import reports

        report = await reports.get_report(session, params.report_id)
        if report.approved_at is None:
            raise BadRequestError("REPORT_NOT_APPROVED", "Rasmiy PDF hisobot tasdiqlangandan keyin tayyorlanadi.")
        return f"hisobot-{str(report.id)[:8]}.pdf", await reports.pdf_bytes(session, report)
    document = await BUILDERS[code](session, params)
    return f"{document.filename}.{fmt}", await document_render.render(document, fmt)
