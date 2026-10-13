"""Round L12–L13: the club's papers. After «Klub rahbari» comes the leader's name (the setting, or the owner —
the first super admin — when it is empty) with the seal stamped over the signature; the ma'lumotnoma says when
it was issued and that it is given to be presented where needed, and keeps its QR code on the same page."""

from __future__ import annotations

import itertools
import re

from sqlalchemy.ext.asyncio import AsyncSession

from app.services import documents, signing
from app.services.document_render import render
from app.services.settings import update_setting
from tests.factories import create_user
from tests.helpers import grant_role

_ids = itertools.count(3_800_000)


async def test_the_leader_signs_and_the_seal_is_stamped(db: AsyncSession) -> None:
    nobody = await signing.signer(db)
    assert nobody.title == "Klub rahbari" and nobody.name == ""  # no super admin yet
    owner = await create_user(db, telegram_id=next(_ids), first_name="Akmaljon", last_name="Bahodirov")
    later = await create_user(db, telegram_id=next(_ids), first_name="Keyingi", last_name="Admin")
    await db.commit()
    await grant_role(db, owner, "SUPER_ADMIN")
    await grant_role(db, later, "SUPER_ADMIN")

    signer = await signing.signer(db)
    assert signer.name == "Akmaljon Bahodirov"  # the owner, without typing anything
    assert signer.seal is not None and signer.seal.startswith("data:image/png;base64,")
    assert signer.line == "Klub rahbari: ____________________ Akmaljon Bahodirov"

    await update_setting(db, key="certificates.signer_name", value="A. Bahodirov", actor_id=owner.id)
    await update_setting(db, key="certificates.seal", value=False, actor_id=owner.id)
    named = await signing.signer(db)
    assert named.name == "A. Bahodirov" and named.seal is None


async def test_the_statement_is_one_complete_page(db: AsyncSession) -> None:
    owner = await create_user(db, telegram_id=next(_ids), first_name="Akmaljon", last_name="Bahodirov")
    member = await create_user(
        db, telegram_id=next(_ids), first_name="Madina", last_name="Karimova", course=3, group_name="5424"
    )
    await db.commit()
    await grant_role(db, owner, "SUPER_ADMIN")

    paper = await documents.statement(db, member)
    assert paper.title == "MA'LUMOTNOMA" and paper.number and paper.number.startswith("№ ")
    assert "Karimova Madinaga" in paper.body[0] and "haqiqatan ham" in paper.body[0]
    assert paper.notes == ["Ushbu ma'lumotnoma kerakli joyga taqdim etish uchun berildi."]
    assert re.fullmatch(r"\d{2}\.\d{2}\.\d{4}, soat \d{2}:\d{2}", paper.issued or "")
    assert paper.inline_qr and paper.signer.name == "Akmaljon Bahodirov"

    from weasyprint import HTML

    from app.services.document_render import _pdf_html
    from app.services.report_pdf import TEMPLATES

    pages = HTML(string=_pdf_html(paper), base_url=str(TEMPLATES)).render().pages
    assert len(pages) == 1  # the text, the figures, the signature with the seal and the QR code — one page
    assert (await render(paper, "pdf")).startswith(b"%PDF")
    sheet = await render(paper, "xlsx")
    from io import BytesIO

    from openpyxl import load_workbook

    cells = [str(cell.value) for row in load_workbook(BytesIO(sheet)).active.iter_rows() for cell in row if cell.value]
    assert any("Klub rahbari: ____" in value and "Bahodirov" in value for value in cells)
    assert any(value.startswith("Berilgan sana va vaqt:") for value in cells)
