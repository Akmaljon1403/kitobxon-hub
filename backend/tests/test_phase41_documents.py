"""Round E11: the document centre (Excel, Word, PDF; downloaded or sent to the bot), the member's own
activity statement, and the portfolio in Word."""

from __future__ import annotations

import io
import itertools
from datetime import UTC, date, datetime, timedelta
from typing import Any

import httpx
import pytest
from aiogram import Bot
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.methods import SendDocument
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import AttendanceRecord, BookRead, Faculty
from app.models.enums import AttendanceMethod, FeatureMode, ReadingStatus, ReportKind
from app.security.permissions import Perm
from app.services import documents, features
from app.services import reports as report_service
from tests.bot_helpers import BOT_TOKEN, FakeTelegramSession
from tests.factories import create_book, create_event, create_test, create_user
from tests.helpers import admin_headers, grant_permissions, grant_role, member_headers
from tests.test_phase25_documents import _attempt

_ids = itertools.count(1_700_000)


@pytest.fixture
def telegram(monkeypatch: pytest.MonkeyPatch) -> FakeTelegramSession:
    session = FakeTelegramSession()
    bot = Bot(token=BOT_TOKEN, session=session, default=DefaultBotProperties(parse_mode=ParseMode.HTML))
    monkeypatch.setattr("app.services.delivery.runtime_bot", lambda: bot)
    return session


async def _club(db: AsyncSession) -> dict[str, Any]:
    faculty = Faculty(name=f"Davolash ishi {next(_ids)}", short_name="DI")
    db.add(faculty)
    await db.flush()
    readers = [
        await create_user(
            db, first_name="Aziza", last_name="Karimova", faculty_id=faculty.id, course=2, group_name="201"
        ),
        await create_user(
            db, first_name="Bobur", last_name="Aliyev", faculty_id=faculty.id, course=3, group_name="301"
        ),
        await create_user(
            db, first_name="Jasur", last_name="Toshev", faculty_id=faculty.id, course=3, group_name="301"
        ),
    ]
    event = await create_event(db, title="O'tkan kunlar muhokamasi", starts_at=datetime.now(UTC) - timedelta(days=2))
    for reader in readers[:2]:
        db.add(AttendanceRecord(event_id=event.id, user_id=reader.id, method=AttendanceMethod.CODE))
    test = await create_test(db, title="Roman testi", event_id=event.id)
    await _attempt(db, test.id, readers[0], 18)
    await _attempt(db, test.id, readers[2], 20)  # did not attend: «davomatsiz»
    book = await create_book(db, title="O'tkan kunlar")
    db.add(
        BookRead(
            book_id=book.id,
            user_id=readers[0].id,
            status=ReadingStatus.COMPLETED,
            completed_at=datetime.now(UTC) - timedelta(days=1),
        )
    )
    await db.commit()
    return {"faculty": faculty, "readers": readers, "event": event, "test": test}


async def _admin(db: AsyncSession, *permissions: Perm) -> Any:
    admin = await create_user(db, telegram_id=next(_ids))
    await db.commit()
    await grant_role(db, admin, "ADMIN")
    await grant_permissions(db, admin, *permissions)
    return admin


def test_study_year_runs_from_september() -> None:
    assert documents.academic_year(date(2026, 10, 10)) == (date(2026, 9, 1), date(2027, 8, 31))
    assert documents.academic_year(date(2027, 3, 1)) == (date(2026, 9, 1), date(2027, 8, 31))


async def test_students_report_in_three_formats(client: httpx.AsyncClient, db: AsyncSession) -> None:
    from docx import Document as WordDocument
    from openpyxl import load_workbook

    club = await _club(db)
    admin = await _admin(db, Perm.ANALYTICS_VIEW)
    headers = admin_headers(admin)
    params = {"faculty_id": str(club["faculty"].id)}

    excel = await client.get(
        "/api/v1/admin/documents/students/file", params={**params, "format": "xlsx"}, headers=headers
    )
    assert excel.status_code == 200, excel.text
    sheet = load_workbook(io.BytesIO(excel.content)).active
    rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    header = next(row for row in rows if row and row[0] == "№")
    names = {row[1]: dict(zip(header, row, strict=False)) for row in rows if row and isinstance(row[0], int)}
    aziza = names["Karimova Aziza"]
    assert aziza["Tadbirlar"] == 1 and aziza["Testlar"] == 1 and aziza["Test ballari"] == 18
    assert aziza["O'qilgan kitoblar"] == 1 and aziza["Kurs"] == 2 and aziza["Guruh"] == "201"
    assert names["Toshev Jasur"]["Tadbirlar"] == 0 and names["Toshev Jasur"]["Testlar"] == 1

    word = await client.get(
        "/api/v1/admin/documents/students/file", params={**params, "format": "docx"}, headers=headers
    )
    text = "\n".join(
        cell.text for table in WordDocument(io.BytesIO(word.content)).tables for row in table.rows for cell in row.cells
    )
    assert "Karimova Aziza" in text and "Aliyev Bobur" in text

    pdf = await client.get("/api/v1/admin/documents/students/file", params={**params, "format": "pdf"}, headers=headers)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert pdf.headers["content-disposition"].startswith("attachment")


async def test_statistics_has_a_qr_appendix(client: httpx.AsyncClient, db: AsyncSession) -> None:
    from docx import Document as WordDocument

    club = await _club(db)
    admin = await _admin(db, Perm.ANALYTICS_VIEW)
    response = await client.get(
        "/api/v1/admin/documents/statistics/file",
        params={"faculty_id": str(club["faculty"].id), "format": "docx"},
        headers=admin_headers(admin),
    )
    assert response.status_code == 200, response.text
    word = WordDocument(io.BytesIO(response.content))
    assert len(word.inline_shapes) == 3  # one QR code per active member
    pdf = await client.get(
        "/api/v1/admin/documents/statistics/file",
        params={"faculty_id": str(club["faculty"].id), "format": "pdf"},
        headers=admin_headers(admin),
    )
    assert pdf.content.startswith(b"%PDF")


async def test_every_template_renders(client: httpx.AsyncClient, db: AsyncSession) -> None:
    club = await _club(db)
    owner = await create_user(db)
    await db.commit()
    await grant_role(db, owner, "SUPER_ADMIN")
    report = await report_service.create_report(db, actor_id=owner.id, kind=ReportKind.EVENT, event_id=club["event"].id)
    await db.commit()
    headers = admin_headers(owner)
    templates = (await client.get("/api/v1/admin/documents", headers=headers)).json()
    assert {item["code"] for item in templates} == set(documents.BY_CODE)
    params = {"event_id": str(club["event"].id), "test_id": str(club["test"].id), "report_id": str(report.id)}
    for template in documents.TEMPLATES:
        for fmt in template.formats:
            response = await client.get(
                f"/api/v1/admin/documents/{template.code}/file", params={**params, "format": fmt}, headers=headers
            )
            if template.code == "event_report" and fmt == "pdf":
                assert response.status_code == 400  # the official PDF exists once the report is approved
                continue
            assert response.status_code == 200, (template.code, fmt, response.text)
            assert len(response.content) > 500
    test_results = await client.get(
        "/api/v1/admin/documents/test_results/file", params={**params, "format": "xlsx"}, headers=headers
    )
    from openpyxl import load_workbook

    values = [cell.value for row in load_workbook(io.BytesIO(test_results.content)).active.iter_rows() for cell in row]
    assert "davomatsiz" in values


async def test_permissions_and_sending_to_the_bot(
    client: httpx.AsyncClient, db: AsyncSession, telegram: FakeTelegramSession
) -> None:
    await _club(db)
    reader = await _admin(db, Perm.USER_VIEW)
    headers = admin_headers(reader)
    listed = {item["code"] for item in (await client.get("/api/v1/admin/documents", headers=headers)).json()}
    assert listed == {"members"}
    denied = await client.get("/api/v1/admin/documents/students/file", params={"format": "xlsx"}, headers=headers)
    assert denied.status_code == 403
    sent = await client.post("/api/v1/admin/documents/members/send", json={"format": "xlsx"}, headers=headers)
    assert sent.status_code == 204, sent.text
    document = next(request for request in telegram.requests if isinstance(request, SendDocument))
    assert document.chat_id == reader.telegram_id and document.document.filename.endswith(".xlsx")
    wrong = await client.post("/api/v1/admin/documents/members/send", json={"format": "pptx"}, headers=headers)
    assert wrong.status_code == 422


async def test_member_statement_and_portfolio_in_word(
    client: httpx.AsyncClient, db: AsyncSession, telegram: FakeTelegramSession
) -> None:
    club = await _club(db)
    member = club["readers"][0]
    owner = await create_user(db)
    await db.commit()
    await features.set_mode(db, actor_id=owner.id, code="portfolio", new_mode=FeatureMode.ALL, announce=False)
    await db.commit()
    headers = member_headers(member)
    for fmt in ("pdf", "docx"):
        response = await client.post("/api/v1/me/statement/send", json={"format": fmt}, headers=headers)
        assert response.status_code == 204, response.text
    names = [request.document.filename for request in telegram.requests if isinstance(request, SendDocument)]
    assert names[0].startswith("malumotnoma-KH-") and names[0].endswith(".pdf") and names[1].endswith(".docx")

    from docx import Document as WordDocument

    statement = next(
        r for r in telegram.requests if isinstance(r, SendDocument) and r.document.filename.endswith(".docx")
    )
    word = WordDocument(io.BytesIO(statement.document.data))
    rows = {row.cells[0].text: row.cells[1].text for row in word.tables[0].rows}
    assert rows["Tadbirlarda qatnashgan"] == "1" and rows["O'qib tugatgan kitoblari"] == "1"
    assert len(word.inline_shapes) == 1  # the QR code

    telegram.requests.clear()
    assert (await client.post("/api/v1/me/portfolio/send", headers=headers)).status_code == 204
    files = [r.document.filename for r in telegram.requests if isinstance(r, SendDocument)]
    assert files[0].endswith(".pdf") and files[1].endswith(".docx")
