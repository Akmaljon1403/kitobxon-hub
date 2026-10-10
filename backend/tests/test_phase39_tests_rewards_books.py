"""Round E6–E8: questions imported from TXT/Word/PDF with pictures, test places with «davomatsiz»,
promo codes handed out by places, and books (covers, reading credited at the event, late reading)."""

from __future__ import annotations

import base64
import io
import itertools
from datetime import UTC, date, datetime, timedelta
from typing import Any

import httpx
import pytest
from PIL import Image
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.models import (
    AttendanceRecord,
    BookRead,
    ChannelPost,
    Notification,
    RewardCode,
    ScoreTransaction,
    TestQuestion,
    UserReward,
)
from app.models.enums import (
    AttendanceMethod,
    ReadingStatus,
    ReadingVerification,
    RewardType,
    ScoreCategory,
    TestStatus,
)
from app.security.permissions import Perm
from app.services import attendance, test_import
from app.services import quizzes as quiz_service
from app.services.storage import get_storage
from tests.factories import create_book, create_event, create_test, create_user
from tests.helpers import admin_headers, grant_permissions, grant_role, member_headers
from tests.test_phase25_documents import _attempt

_ids = itertools.count(1_500_000)


def _png(color: tuple[int, int, int] = (200, 30, 30)) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (320, 200), color).save(buffer, format="PNG")
    return buffer.getvalue()


async def _editor(db: AsyncSession, *roles: str) -> Any:
    admin = await create_user(db, telegram_id=next(_ids))
    await db.commit()
    await grant_role(db, admin, "ADMIN")
    for role in roles:
        await grant_role(db, admin, role)
    await grant_permissions(
        db,
        admin,
        Perm.TEST_VIEW,
        Perm.TEST_EDIT,
        Perm.TEST_PUBLISH,
        Perm.BOOK_VIEW,
        Perm.BOOK_EDIT,
        Perm.REWARD_MANAGE,
        Perm.EVENT_VIEW,
        Perm.ATTENDANCE_MANAGE,
    )
    return admin


# ---------------------------------------------------------------------------- import


TXT = """Kitob testi bo'yicha savollar
1. «O'tkan kunlar» romani muallifi kim? (2 ball)
A) Cho'lpon
*B) Abdulla Qodiriy
C) Oybek
Izoh: Roman 1922–1925-yillarda yozilgan.

2) Qaysilari shoir?
A. Navoiy
B. Bobur
C. Qodiriy
Javob: A, B

3. Bu savolda javob yo'q
A) bir
B) ikki
"""


def test_text_questions_are_understood() -> None:
    questions, warnings = test_import.to_questions(test_import.parse(test_import.read("t.txt", TXT.encode())))
    assert [q["type"] for q in questions] == ["SINGLE_CHOICE", "MULTIPLE_CHOICE", "SINGLE_CHOICE"]
    first = questions[0]
    assert first["points"] == 2 and first["text"] == "«O'tkan kunlar» romani muallifi kim?"
    assert [o["is_correct"] for o in first["options"]] == [False, True, False]
    assert first["explanation"] == "Roman 1922–1925-yillarda yozilgan."
    assert [o["is_correct"] for o in questions[1]["options"]] == [True, True, False]
    assert warnings == ["3-savol: to'g'ri javob belgilanmagan"]
    # Plus/minus style and a Windows-1251 file.
    signed = "? Kitob nomi\n+ Mehrobdan chayon\n- Sarob\n".encode("cp1251")
    questions, _ = test_import.to_questions(test_import.parse(test_import.read("t.txt", signed)))
    assert [(o["text"], o["is_correct"]) for o in questions[0]["options"]] == [
        ("Mehrobdan chayon", True),
        ("Sarob", False),
    ]


async def test_word_and_pdf_imports_keep_pictures(client: httpx.AsyncClient, db: AsyncSession) -> None:
    from docx import Document
    from docx.shared import Inches

    admin = await _editor(db)
    test = await create_test(db, status=TestStatus.DRAFT, starts_at=datetime.now(UTC) + timedelta(days=1))
    await db.commit()
    headers = admin_headers(admin)

    document = Document()
    document.add_paragraph("1. Rasmda kim tasvirlangan?")
    document.add_picture(io.BytesIO(_png()), width=Inches(2))
    document.add_paragraph("A) Navoiy")
    run = document.add_paragraph().add_run("B) Bobur")
    run.bold = True  # the answer marked in bold
    document.add_paragraph("C) Ogahiy")
    buffer = io.BytesIO()
    document.save(buffer)
    response = await client.post(
        f"/api/v1/admin/tests/{test.id}/import",
        files={"file": ("savollar.docx", buffer.getvalue(), "application/octet-stream")},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    question = response.json()["questions"][0]
    assert [o["is_correct"] for o in question["options"]] == [False, True, False]
    assert question["image_key"].startswith(f"tests/{test.id}/")
    picture = await client.get(
        f"/api/v1/admin/tests/{test.id}/question-images", params={"key": question["image_key"]}, headers=headers
    )
    assert picture.status_code == 200 and picture.content[:2] == b"\xff\xd8"  # stored as JPEG

    from weasyprint import HTML

    uri = "data:image/png;base64," + base64.b64encode(_png((10, 120, 10))).decode()
    html = f"<p>1. Birinchi savol?</p><img src='{uri}' style='width:200px'><p>A) bir</p><p>*B) ikki</p>"
    pdf = await client.post(
        f"/api/v1/admin/tests/{test.id}/import",
        files={"file": ("savollar.pdf", HTML(string=html).write_pdf(), "application/pdf")},
        headers=headers,
    )
    assert pdf.status_code == 200, pdf.text
    assert pdf.json()["questions"][0]["image_key"] is not None

    empty = await client.post(
        f"/api/v1/admin/tests/{test.id}/import", files={"file": ("x.txt", b"hech narsa", "text/plain")}, headers=headers
    )
    assert empty.status_code == 400 and empty.json()["error"]["code"] == "IMPORT_NO_QUESTIONS"


async def test_question_pictures_are_saved_and_shown_to_members(client: httpx.AsyncClient, db: AsyncSession) -> None:
    admin = await _editor(db)
    test = await create_test(db, status=TestStatus.DRAFT)
    member = await create_user(db)
    await db.commit()
    headers = admin_headers(admin)
    uploaded = await client.post(
        f"/api/v1/admin/tests/{test.id}/question-images",
        files={"file": ("rasm.png", _png(), "image/png")},
        headers=headers,
    )
    key = uploaded.json()["image_key"]
    question = {
        "type": "SINGLE_CHOICE",
        "text": "Rasmdagi joy?",
        "points": 1,
        "options": [{"text": "Samarqand", "is_correct": True}, {"text": "Buxoro", "is_correct": False}],
    }
    foreign = await client.put(
        f"/api/v1/admin/tests/{test.id}/questions",
        json={"questions": [{**question, "image_key": "tests/other/q.jpg"}]},
        headers=headers,
    )
    assert foreign.status_code == 400
    saved = await client.put(
        f"/api/v1/admin/tests/{test.id}/questions",
        json={"questions": [{**question, "image_key": key}]},
        headers=headers,
    )
    assert saved.status_code == 200 and saved.json()["questions"][0]["image_key"] == key
    assert (await client.post(f"/api/v1/admin/tests/{test.id}/publish", headers=headers)).status_code == 200
    question_id = saved.json()["questions"][0]["id"]

    image_url = f"/api/v1/tests/{test.id}/questions/{question_id}/image"
    before = await client.get(image_url, headers=member_headers(member))
    assert before.status_code == 403  # not before the member starts the test
    started = await client.post(f"/api/v1/tests/{test.id}/start", headers=member_headers(member))
    assert started.status_code in (200, 201), started.text
    assert started.json()["questions"][0]["has_image"] is True
    after = await client.get(image_url, headers=member_headers(member))
    assert after.status_code == 200, after.text and after.content[:2] == b"\xff\xd8"


async def test_removed_question_takes_its_picture_away(client: httpx.AsyncClient, db: AsyncSession) -> None:
    admin = await _editor(db)
    test = await create_test(db, status=TestStatus.DRAFT)
    await db.commit()
    headers = admin_headers(admin)
    key = (
        await client.post(
            f"/api/v1/admin/tests/{test.id}/question-images",
            files={"file": ("r.png", _png(), "image/png")},
            headers=headers,
        )
    ).json()["image_key"]
    base = {"type": "SINGLE_CHOICE", "points": 1, "options": [{"text": "a", "is_correct": True}, {"text": "b"}]}
    await client.put(
        f"/api/v1/admin/tests/{test.id}/questions",
        json={"questions": [{**base, "text": "Bir", "image_key": key}]},
        headers=headers,
    )
    assert await get_storage().get(key)
    await client.put(
        f"/api/v1/admin/tests/{test.id}/questions", json={"questions": [{**base, "text": "Ikki"}]}, headers=headers
    )
    from app.core.errors import AppError

    with pytest.raises(AppError):
        await get_storage().get(key)
    assert (
        await db.scalar(select(func.count()).select_from(TestQuestion).where(TestQuestion.image_key.is_not(None))) == 0
    )


# ---------------------------------------------------------------------------- places with «davomatsiz»


async def _event_test(db: AsyncSession, scores: list[int], attended: list[bool]) -> tuple[Any, list[Any]]:
    event = await create_event(db, title="Kitob kechasi")
    test = await create_test(db, title="Kecha testi", event_id=event.id)
    members = [await create_user(db, first_name=f"Ishtirokchi{i}") for i in range(len(scores))]
    for member, score, was_there in zip(members, scores, attended, strict=True):
        await _attempt(db, test.id, member, score)
        if was_there:
            db.add(AttendanceRecord(event_id=event.id, user_id=member.id, method=AttendanceMethod.CODE))
    await db.commit()
    return test, members


async def test_non_attendees_are_listed_but_take_no_prize_places(
    client: httpx.AsyncClient, db: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    test, members = await _event_test(db, [20, 18, 15, 12], [False, True, True, True])
    test_id = test.id
    await quiz_service.close_test(db, actor_id=None, test_id=test_id)

    rows = await quiz_service.standings_with_attendance(db, await quiz_service.get_test(db, test_id))
    assert [(r["place"], r["attended"], r["official_place"]) for r in rows] == [
        (1, False, None),
        (2, True, 1),
        (3, True, 2),
        (4, True, 3),
    ]
    # Bonus points only for the places that count.
    bonus = list(
        await db.scalars(
            select(ScoreTransaction.user_id).where(
                ScoreTransaction.category == ScoreCategory.TEST_RANK_BONUS, ScoreTransaction.source_id == test_id
            )
        )
    )
    assert members[0].id not in bonus and members[1].id in bonus
    # The best result without attendance gets a kind note instead of a place.
    note = await db.scalar(
        select(Notification.text).where(Notification.dedup_key == f"test_result:{test_id}:{members[0].id}")
    )
    assert "davomat qilinmagani" in note
    winner = await db.scalar(
        select(Notification.text).where(Notification.dedup_key == f"test_result:{test_id}:{members[1].id}")
    )
    assert "1-o'rinni" in winner
    post = await db.scalar(select(ChannelPost.text).where(ChannelPost.dedup_key == f"test:{test_id}"))
    assert post is None or "Ishtirokchi0" not in post

    table = await client.get(f"/api/v1/tests/{test_id}/standings", headers=member_headers(members[0]))
    assert table.status_code == 200, table.text
    data = table.json()
    assert data["event_test"] is True
    assert data["rows"][0]["attended"] is False and data["rows"][0]["prize_place"] is None and data["rows"][0]["me"]
    assert data["rows"][1]["prize_place"] == 1

    admin = await _editor(db)
    results = (await client.get(f"/api/v1/admin/tests/{test_id}/results", headers=admin_headers(admin))).json()
    first = next(item for item in results["items"] if item["user_id"] == str(members[0].id))
    assert first["attended"] is False and first["prize_place"] is None and first["place"] == 1


# ---------------------------------------------------------------------------- promo codes by places


async def test_codes_go_to_first_places_first_one_per_member(client: httpx.AsyncClient, db: AsyncSession) -> None:
    admin = await _editor(db)
    one, members_one = await _event_test(db, [20, 18, 15], [True, True, True])
    two_members = [
        members_one[0],
        await create_user(db, first_name="Boshqa"),
        await create_user(db, first_name="Uchinchi"),
    ]
    two = await create_test(db, title="Ikkinchi test")
    for member, score in zip(two_members, (19, 17, 11), strict=True):
        await _attempt(db, two.id, member, score)
    await db.commit()
    one_id, two_id = one.id, two.id
    await quiz_service.close_test(db, actor_id=None, test_id=one_id)
    await quiz_service.close_test(db, actor_id=None, test_id=two_id)
    from app.services import rewards

    reward = await rewards.create_reward(
        db, actor_id=admin.id, data={"title": "Asaxiy chegirma", "type": RewardType.PARTNER_DISCOUNT}
    )
    reward_id = reward.id
    headers = admin_headers(admin)

    plan = (await client.get("/api/v1/admin/rewards/distribution-plan", headers=headers)).json()["waiting"]
    # 1st places first (the same member won both tests: one code only), then the 2nd places, then the 3rd.
    assert [slot["place"] for slot in plan] == [1, 2, 2, 3, 3]
    assert len({slot["user_id"] for slot in plan}) == len(plan)

    response = await client.post(
        f"/api/v1/admin/rewards/{reward_id}/distribute-test-places",
        json={"codes": "KOD-1\nKOD-2\n\nKOD-3\nKOD-1", "usage": "Asaxiy.uz saytida to'lovda kiriting", "days": 90},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["codes_added"] == 3  # the repeated KOD-1 is counted once
    assert [slot["place"] for slot in result["given"]] == [1, 2, 2]
    assert [slot["place"] for slot in result["waiting"]] == [3, 3]
    message = await db.scalar(
        select(Notification.text).where(Notification.user_id == members_one[0].id, Notification.type == "REWARD")
    )
    assert "Asaxiy.uz saytida" in message and "1-o'rinni" in message
    assert "KOD-" not in message  # the code itself is seen only in the app

    # More codes later: only the places still waiting get them.
    again = await client.post(
        f"/api/v1/admin/rewards/{reward_id}/distribute-test-places",
        json={"codes": "KOD-4\nKOD-5\nKOD-6", "usage": None, "days": 90},
        headers=headers,
    )
    assert [slot["place"] for slot in again.json()["given"]] == [3, 3] and again.json()["waiting"] == []
    assert await db.scalar(select(func.count()).select_from(UserReward).where(UserReward.reward_id == reward_id)) == 5
    left = await db.scalar(
        select(func.count())
        .select_from(RewardCode)
        .where(RewardCode.reward_id == reward_id, RewardCode.status == "AVAILABLE")
    )
    assert left == 1


# ---------------------------------------------------------------------------- books


async def test_book_cover_upload_is_public(client: httpx.AsyncClient, db: AsyncSession) -> None:
    admin = await _editor(db)
    book = await create_book(db, title="Kecha va kunduz")
    member = await create_user(db)
    await db.commit()
    response = await client.put(
        f"/api/v1/admin/books/{book.id}/cover",
        files={"file": ("cover.png", _png(), "image/png")},
        headers=admin_headers(admin),
    )
    assert response.status_code == 200, response.text
    cover = response.json()["cover_image"]
    assert cover.startswith(f"/api/v1/meta/book-covers/{book.id}.jpg?v=") and response.json()["cover_uploaded"]
    picture = await client.get(cover)
    assert picture.status_code == 200 and picture.content[:2] == b"\xff\xd8"
    from app.models.enums import BookStatus

    book.status = BookStatus.ACTIVE
    await db.commit()
    listed = (await client.get(f"/api/v1/books/{book.id}", headers=member_headers(member))).json()
    assert listed["cover_url"] == cover
    removed = await client.delete(f"/api/v1/admin/books/{book.id}/cover", headers=admin_headers(admin))
    assert removed.json()["cover_uploaded"] is False
    assert (await client.get(cover)).status_code == 404


async def test_attending_the_books_event_counts_as_reading(db: AsyncSession) -> None:
    owner = await _editor(db, "SUPER_ADMIN")
    book = await create_book(db, title="Sarob")
    event = await create_event(db, title="Sarob muhokamasi", book_id=book.id)
    member = await create_user(db)
    await db.commit()
    await attendance.manual_check_in(db, actor_id=owner.id, event_id=event.id, user_id=member.id, note="Kod ishlamadi")
    read = await db.scalar(select(BookRead).where(BookRead.user_id == member.id, BookRead.book_id == book.id))
    assert read.status == ReadingStatus.COMPLETED and read.verification_source == ReadingVerification.EVENT
    points = await db.scalar(
        select(func.count())
        .select_from(ScoreTransaction)
        .where(ScoreTransaction.user_id == member.id, ScoreTransaction.category == ScoreCategory.BOOK_COMPLETION)
    )
    assert points == 1


async def test_reading_after_the_deadline_has_no_points(
    client: httpx.AsyncClient, db: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.models.enums import BookStatus
    from app.services.settings import update_setting

    member = await create_user(db)
    late = await create_book(db, status=BookStatus.COMPLETED, deadline=date.today() - timedelta(days=3))
    await update_setting(db, key="score.book_completion_requires_verification", value=False, actor_id=member.id)
    await db.commit()
    headers = member_headers(member)
    assert (
        await client.post(f"/api/v1/books/{late.id}/reading", json={"action": "start"}, headers=headers)
    ).status_code == 200
    done = await client.post(f"/api/v1/books/{late.id}/reading", json={"action": "complete"}, headers=headers)
    assert done.status_code == 200 and done.json()["my_reading"]["status"] == "COMPLETED"
    assert done.json()["my_reading"]["late"] is True
    points = await db.scalar(
        select(func.count())
        .select_from(ScoreTransaction)
        .where(ScoreTransaction.user_id == member.id, ScoreTransaction.category == ScoreCategory.BOOK_COMPLETION)
    )
    assert points == 0
    assert get_settings() is not None
